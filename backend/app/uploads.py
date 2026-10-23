"""Загрузка файлов: распознать → предпросмотр → «Применить» → при необходимости откат.

Поддерживаются:
  • шаблон FiberTrack (его выдаёт сама программа) — пробег, выручка и расходы по машинам, общие расходы;
  • таблица прибыли по машинам (формат заказчика) — только владелец;
  • файл парка: график работы, «Анализ ЗАПРАВОК», «АВТО» (любой набор этих листов).

Правила:
  • закрытые месяцы не меняются никогда — строки по ним показываются как пропущенные;
  • ничего не угадываем: непонятные машины и значения — ошибка в предпросмотре и пункт в «Проверить»;
  • «Применить» выполняет ровно тот план, что показан в предпросмотре;
  • откат возвращает данные как были; если их после загрузки меняли, откат отказывается и говорит почему.
"""
from __future__ import annotations

import datetime as dt
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy import delete, inspect, select
from sqlalchemy.orm import Session

from . import parsing
from .audit import log
from .data import days_in, ui_settings
from .models import (FuelDay, Mileage, Operation, Period, ReferenceValue, ReviewItem, ScheduleMonth, Upload, User, Vehicle)
from .security import now

MONTHS = ['январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь']
TEMPLATE_MARK = 'fibertrack-template'
# столбцы шаблона: (код статьи, заголовок). km — пробег, остальное — деньги
TEMPLATE_COLUMNS = [
    ('km', 'Пробег, км'), ('revenue', 'Выручка от перевозок'), ('other_income', 'Прочий доход'),
    ('fuel', 'Топливо'), ('repair', 'Ремонт и запчасти'), ('maintenance', 'ТО'), ('tires', 'Шины'),
    ('salary', 'Зарплата водителей'), ('per_diem', 'Суточные'), ('tolls', 'Платные дороги и «Платон»'),
    ('insurance', 'Страхование'), ('leasing', 'Лизинг / кредит'), ('wash', 'Мойка'), ('fines', 'Штрафы'), ('other', 'Прочее'),
]
MONEY_ARTICLES = [c for c, _ in TEMPLATE_COLUMNS if c != 'km']
INCOME_ARTICLES = ('revenue', 'other_income')
SHARED_ROWS = {'ОБЩЕХОЗ': 'shared_gen', 'СЕРВИС': 'shared_service', 'ТРАКТОР': 'shared_tractor'}
VAT_RATES = {'22%': 22, '10%': 10, '7%': 7, '5%': 5, '0%': 0, 'без НДС': None}
TRACKED = (Operation, Mileage, ScheduleMonth, FuelDay)


class UploadError(Exception):
    """Ошибка, понятная пользователю."""


def ym_label(ym: str) -> str:
    y, m = ym.split('-')
    return f'{MONTHS[int(m) - 1].capitalize()} {y}'


def label_to_ym(text: str) -> str | None:
    m = re.search(r'(январ|феврал|март|апрел|ма[йя]|июн|июл|август|сентябр|октябр|ноябр|декабр)\w*\s+(20\d\d)', str(text or '').lower())
    if not m:
        return None
    stems = ['январ', 'феврал', 'март', 'апрел', 'ма', 'июн', 'июл', 'август', 'сентябр', 'октябр', 'ноябр', 'декабр']
    stem = m.group(1)
    idx = next(i for i, s in enumerate(stems) if stem.startswith(s))
    return f'{m.group(2)}-{idx + 1:02d}'


def plural(n: int, one: str, few: str, many: str) -> str:
    """3 машины, 5 машин, 21 машина."""
    n10, n100 = n % 10, n % 100
    w = one if n10 == 1 and n100 != 11 else few if 2 <= n10 <= 4 and not 12 <= n100 <= 14 else many
    return f'{n} {w}'


def rub(kop: int) -> str:
    v = kop / 100
    s = f'{abs(v):,.0f}'.replace(',', ' ')
    return ('−' if v < 0 else '') + s + ' ₽'


# ---------------------------------------------------------------- значения ячеек
def money_kop(v) -> int | None:
    """Сумма из ячейки → копейки. Пусто → None. Не число → ValueError."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ValueError
    if isinstance(v, (int, float, Decimal)):
        return int((Decimal(str(v)) * 100).quantize(Decimal('1')))
    s = str(v).replace(' ', '').replace(' ', '').replace('₽', '').replace('руб.', '').replace('руб', '').replace(',', '.').replace('−', '-')
    try:
        return int((Decimal(s) * 100).quantize(Decimal('1')))
    except InvalidOperation:
        raise ValueError


def int_value(v) -> int | None:
    k = money_kop(v)
    if k is None:
        return None
    if k % 100:
        raise ValueError
    return k // 100


def vat_value(v) -> str | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        p = v * 100 if 0 < v < 1 else v
        key = f'{int(round(p))}%'
        if key in VAT_RATES:
            return key
        raise ValueError
    s = str(v).strip().lower().replace(' ', '')
    if s in ('безндс', 'бездс', 'нет', '-', '—'):
        return 'без НДС'
    s = s.replace(',', '.').rstrip('%')
    try:
        p = float(s)
    except ValueError:
        raise ValueError
    key = f'{int(round(p * 100 if 0 < p < 1 else p))}%'
    if key in VAT_RATES:
        return key
    raise ValueError


def yes_no(v) -> bool | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in ('да', 'yes', 'true', '1', 'с ндс', '+'):
        return True
    if s in ('нет', 'no', 'false', '0', 'без ндс', '-'):
        return False
    raise ValueError


def split_vat(amount_kop: int, rate: int | None, gross: bool) -> tuple[int, int, int]:
    """→ (брутто, нетто, НДС) в копейках."""
    if not rate:
        return amount_kop, amount_kop, 0
    if gross:
        net = round(amount_kop * 100 / (100 + rate))
        return amount_kop, net, amount_kop - net
    vat = round(amount_kop * rate / 100)
    return amount_kop + vat, amount_kop, vat


# ---------------------------------------------------------------- снимки строк для отката
def row_to_dict(obj) -> dict:
    out = {}
    for col in inspect(obj).mapper.column_attrs:
        v = getattr(obj, col.key)
        if isinstance(v, (dt.date, dt.datetime)):
            v = {'__date__' if type(v) is dt.date else '__dt__': v.isoformat()}
        elif isinstance(v, Decimal):
            v = {'__dec__': str(v)}
        out[col.key] = v
    return out


def dict_to_row(model, d: dict):
    vals = {}
    for k, v in d.items():
        if isinstance(v, dict) and '__date__' in v:
            v = dt.date.fromisoformat(v['__date__'])
        elif isinstance(v, dict) and '__dt__' in v:
            v = dt.datetime.fromisoformat(v['__dt__'])
        elif isinstance(v, dict) and '__dec__' in v:
            v = Decimal(v['__dec__'])
        vals[k] = v
    return model(**vals)


# ---------------------------------------------------------------- общий контекст
class Ctx:
    def __init__(self, db: Session):
        self.db = db
        self.vehicles = {v.key: v for v in db.scalars(select(Vehicle))}
        self.by_plate = {v.plate: v for v in self.vehicles.values() if v.plate}
        self.periods = {p.ym: p.status for p in db.scalars(select(Period))}
        self.settings = ui_settings(db)

    def is_open(self, ym: str) -> bool:
        return self.periods.get(ym) == 'open'

    def find_vehicle(self, raw) -> Vehicle | None:
        if raw is None:
            return None
        plate = re.sub(r'\s+', '', str(raw).upper()).translate(str.maketrans('ABEKMHOPCTYX', 'АВЕКМНОРСТУХ'))
        if plate in self.by_plate:
            return self.by_plate[plate]
        key = parsing.car_key(raw)
        return self.vehicles.get(key) if key else None


def section(key, title, columns, rows, notes=None, summary=''):
    errors = sum(1 for r in rows if r['status'] == 'error')
    applied = sum(1 for r in rows if r['status'] in ('ok', 'warn'))
    return {'key': key, 'title': title, 'columns': columns, 'rows': rows, 'notes': notes or [], 'summary': summary,
            'stats': {'apply': applied, 'errors': errors, 'skip': sum(1 for r in rows if r['status'] == 'skip')}}


def row(cells, status='ok', note=''):
    return {'cells': cells, 'status': status, 'note': note}


# ================================================================ распознавание
def detect(wb) -> str:
    names = [ws.title for ws in wb.worksheets]
    if '_ft' in names and str(wb['_ft']['A1'].value or '') == TEMPLATE_MARK:
        return 'template'
    for ws in wb.worksheets[:3]:
        a = str(ws['A1'].value or '')
        if a.startswith('FiberTrack · шаблон') or (str(ws['A2'].value or '').strip() == 'Машина' and str(ws['B2'].value or '').strip() == 'Госномер'):
            return 'template'
    ws, _ = parsing.find_profit_sheet(wb)
    if ws is not None:
        return 'profit'
    if 'АВТО' in names or 'Анализ ЗАПРАВОК' in names or parsing.parse_schedule(wb, only_check=True):
        return 'park'
    raise UploadError('Не узнал формат файла. Подходят: шаблон FiberTrack (кнопка «Скачать шаблон на месяц»), '
                      'таблица прибыли по машинам, файл парка с графиком работы, «Анализом ЗАПРАВОК» или листом «АВТО».')


# ================================================================ шаблон FiberTrack
def build_template(db: Session, ym: str, keys: list[str]) -> openpyxl.Workbook:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ym_label(ym)
    navy, yellow = PatternFill('solid', fgColor='0B1F3A'), PatternFill('solid', fgColor='FFF3CC')
    thin = Side(style='thin', color='D6D9DF')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = ['Машина', 'Госномер'] + [t for _, t in TEMPLATE_COLUMNS] + ['Ставка НДС', 'Суммы с НДС (да/нет)', 'Документ']
    ws['A1'] = f'FiberTrack · шаблон ввода · {ym_label(ym)} · заполните жёлтые ячейки, суммы в рублях'
    ws['A1'].font = Font(bold=True, size=13)
    for c, h in enumerate(head, 1):
        cell = ws.cell(2, c, h)
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), navy
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 34
    vehicles = {v.key: v for v in db.scalars(select(Vehicle).where(Vehicle.key.in_(keys)))}
    order = sorted(keys, key=lambda k: (int(re.match(r'\d+', k).group()), k))
    r = 3
    for k in order:
        v = vehicles[k]
        ws.cell(r, 1, v.name).font = Font(bold=True)
        ws.cell(r, 2, v.plate or '')
        for c in range(3, 3 + len(TEMPLATE_COLUMNS)):
            cell = ws.cell(r, c)
            cell.fill, cell.border, cell.number_format = yellow, box, '#,##0.00;[Red]-#,##0.00' if c > 3 else '#,##0'
        ws.cell(r, 3 + len(TEMPLATE_COLUMNS), '22%').fill = yellow
        ws.cell(r, 4 + len(TEMPLATE_COLUMNS), 'да').fill = yellow
        ws.cell(r, 5 + len(TEMPLATE_COLUMNS)).fill = yellow
        r += 1
    last_car = r - 1
    r += 1
    for c, h in enumerate(['Общие расходы', 'Сумма', 'Ставка НДС', 'Суммы с НДС (да/нет)', 'Документ'], 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), navy
    for name in ('Общехоз', 'Сервис', 'Трактор'):
        r += 1
        ws.cell(r, 1, name).font = Font(bold=True)
        ws.cell(r, 2).fill, ws.cell(r, 2).border, ws.cell(r, 2).number_format = yellow, box, '#,##0.00;[Red]-#,##0.00'
        ws.cell(r, 3, '22%').fill = yellow
        ws.cell(r, 4, 'да').fill = yellow
        ws.cell(r, 5).fill = yellow
    vat_col = openpyxl.utils.get_column_letter(3 + len(TEMPLATE_COLUMNS))
    gross_col = openpyxl.utils.get_column_letter(4 + len(TEMPLATE_COLUMNS))
    dv_vat = DataValidation(type='list', formula1='"' + ','.join(VAT_RATES) + '"', allow_blank=False, showErrorMessage=True,
                            errorTitle='Ставка НДС', error='Выберите ставку из списка')
    dv_yes = DataValidation(type='list', formula1='"да,нет"', allow_blank=False)
    ws.add_data_validation(dv_vat)
    ws.add_data_validation(dv_yes)
    if last_car >= 3:
        dv_vat.add(f'{vat_col}3:{vat_col}{last_car}')
        dv_yes.add(f'{gross_col}3:{gross_col}{last_car}')
    dv_vat.add(f'C{last_car + 3}:C{r}')
    dv_yes.add(f'D{last_car + 3}:D{r}')
    widths = [12, 13] + [14] * len(TEMPLATE_COLUMNS) + [12, 12, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'C3'
    meta = wb.create_sheet('_ft')
    meta['A1'], meta['A2'] = TEMPLATE_MARK, ym
    meta.sheet_state = 'hidden'
    return wb


def plan_template(ctx: Ctx, wb, chosen_ym: str | None) -> dict:
    ws = wb.worksheets[0]
    ym = None
    if '_ft' in wb.sheetnames:
        ym = str(wb['_ft']['A2'].value or '') or None
    ym = ym or label_to_ym(ws['A1'].value) or label_to_ym(ws.title) or chosen_ym
    if not ym:
        raise UploadError('В шаблоне не указан месяц. Скачайте шаблон заново кнопкой «Скачать шаблон на месяц».')
    g = parsing.grid(ws)
    hdr_i = next((i for i, r in enumerate(g) if str(r[0] or '').strip() == 'Машина'), None)
    if hdr_i is None:
        raise UploadError('В шаблоне не найдена строка заголовков («Машина», «Госномер», …). Скачайте шаблон заново.')
    hdr = [str(x or '').strip() for x in g[hdr_i]]
    col = {}
    for code, title in TEMPLATE_COLUMNS:
        if title in hdr:
            col[code] = hdr.index(title)
    for code, title in (('vat', 'Ставка НДС'), ('gross', 'Суммы с НДС (да/нет)'), ('doc', 'Документ')):
        if title in hdr:
            col[code] = hdr.index(title)
    missing = [t for c, t in TEMPLATE_COLUMNS if c not in col]
    if missing:
        raise UploadError('В шаблоне не хватает столбцов: ' + ', '.join(missing) + '. Скачайте шаблон заново.')
    s = ctx.settings
    month_open = ctx.is_open(ym)
    rows, cars, shared = [], [], []
    seen = set()
    i = hdr_i + 1
    while i < len(g) and str(g[i][0] or '').strip() != 'Общие расходы':
        r = g[i]
        i += 1
        name = r[0]
        if name is None or not str(name).strip():
            continue
        vals = {c: r[col[c]] for c in col}
        filled = any(vals[c] not in (None, '') for c, _ in TEMPLATE_COLUMNS)
        v = ctx.find_vehicle(name) or ctx.find_vehicle(r[1])
        cells = {'car': str(name).strip(), 'plate': str(r[1] or '')}
        if not filled:
            rows.append(row([cells['car'], cells['plate'], '', '', '', ''], 'skip', 'пусто — не меняется'))
            continue
        errs = []
        if v is None:
            errs.append('машина не найдена в справочнике')
        elif v.key in seen:
            errs.append('машина повторяется в шаблоне')
        try:
            km = int_value(vals['km'])
        except ValueError:
            km, _ = None, errs.append('пробег — не целое число')
        if km is not None and not (s['kmMin'] <= km <= s['kmMax']):
            errs.append(f"пробег вне пределов {int(s['kmMin'])}–{int(s['kmMax'])} км")
        amounts = {}
        for a in MONEY_ARTICLES:
            try:
                k = money_kop(vals[a])
            except ValueError:
                errs.append(f'«{dict(TEMPLATE_COLUMNS)[a]}» — не число')
                continue
            if k is not None:
                if k < 0:
                    errs.append(f'«{dict(TEMPLATE_COLUMNS)[a]}» отрицательная')
                amounts[a] = k
        try:
            vat = vat_value(vals.get('vat'))
        except ValueError:
            vat, _ = None, errs.append('непонятная ставка НДС')
        try:
            gross = yes_no(vals.get('gross'))
        except ValueError:
            gross, _ = None, errs.append('«Суммы с НДС» — укажите да или нет')
        if amounts and vat is None and not any('НДС' in e for e in errs):
            errs.append('не указана ставка НДС')
        gross = True if gross is None else gross
        if vat == 'без НДС' and gross and any(amounts.values()):
            errs.append('ставка «без НДС», а суммы отмечены «с НДС»')
        income = sum(amounts.get(a, 0) for a in INCOME_ARTICLES)
        expense = sum(k for a, k in amounts.items() if a not in INCOME_ARTICLES)
        disp = [v.name if v else cells['car'], (v.plate if v else cells['plate']) or '', '' if km is None else f'{km:,}'.replace(',', ' '),
                rub(income) if income else '', rub(expense) if expense else '', vat or '']
        if errs:
            rows.append(row(disp, 'error', '; '.join(errs)))
            continue
        if not month_open:
            rows.append(row(disp, 'skip', 'месяц закрыт'))
            continue
        seen.add(v.key)
        cars.append({'vehicle_id': v.id, 'key': v.key, 'km': km, 'amounts': amounts, 'vat': vat or '22%', 'gross': gross,
                     'doc': str(vals.get('doc') or '').strip()[:128] or None})
        rows.append(row(disp, 'ok', 'будет записано'))
    # общие расходы
    shared_rows = []
    while i < len(g):
        r = g[i]
        i += 1
        name = str(r[0] or '').strip()
        art = SHARED_ROWS.get(name.upper())
        if not art:
            continue
        errs = []
        try:
            k = money_kop(r[1])
        except ValueError:
            k, _ = None, errs.append('сумма — не число')
        if k is None and not errs:
            shared_rows.append(row([name, '', '', ''], 'skip', 'пусто — не меняется'))
            continue
        try:
            vat = vat_value(r[2]) or None
        except ValueError:
            vat, _ = None, errs.append('непонятная ставка НДС')
        try:
            gross = yes_no(r[3])
        except ValueError:
            gross, _ = None, errs.append('«Суммы с НДС» — укажите да или нет')
        gross = True if gross is None else gross
        if vat is None and not errs:
            errs.append('не указана ставка НДС')
        if k is not None and k < 0:
            errs.append('сумма отрицательная')
        disp = [name, rub(k) if k is not None else '', vat or '', 'да' if gross else 'нет']
        if errs:
            shared_rows.append(row(disp, 'error', '; '.join(errs)))
        elif not month_open:
            shared_rows.append(row(disp, 'skip', 'месяц закрыт'))
        else:
            shared.append({'article': art, 'kop': k, 'vat': vat, 'gross': gross, 'doc': str(r[4] or '').strip()[:128] or None})
            shared_rows.append(row(disp, 'ok', 'будет записано'))
    notes = []
    if not month_open:
        notes.append(f'{ym_label(ym)} {"закрыт" if ctx.periods.get(ym) == "closed" else "не открыт"} для ввода — ничего не будет записано. '
                     'Попросите владельца открыть месяц.')
    if chosen_ym and chosen_ym != ym:
        notes.append(f'Шаблон за {ym_label(ym).lower()}, а на экране выбран {ym_label(chosen_ym).lower()}. Данные пойдут в {ym_label(ym).lower()} — как в шаблоне.')
    sections = [section('cars', f'Машины · {ym_label(ym)}', ['Машина', 'Госномер', 'Пробег', 'Доходы', 'Расходы', 'НДС'], rows, notes,
                        plural(len(cars), 'машина', 'машины', 'машин') + ' к записи')]
    if shared_rows:
        sections.append(section('shared', 'Общие расходы', ['Статья', 'Сумма', 'НДС', 'С НДС'], shared_rows))
    return {'kind': 'template', 'months': [ym] if (cars or shared) else [], 'ym': ym, 'sections': sections,
            'plan': {'ym': ym, 'cars': cars, 'shared': shared}}


def apply_template(ctx: Ctx, up: Upload, user: User, plan: dict, backup: dict) -> list[str]:
    db, ym = ctx.db, plan['ym']
    for c in plan['cars']:
        vid = c['vehicle_id']
        old_ops = list(db.scalars(select(Operation).where(Operation.ym == ym, Operation.vehicle_id == vid, Operation.source.in_(('manual', 'template')))))
        old_mil = list(db.scalars(select(Mileage).where(Mileage.ym == ym, Mileage.vehicle_id == vid)))
        backup['operations'] += [row_to_dict(o) for o in old_ops]
        backup['mileage'] += [row_to_dict(m) for m in old_mil]
        for o in old_ops + old_mil:
            db.delete(o)
        db.flush()
        if c['km'] is not None:
            db.add(Mileage(ym=ym, vehicle_id=vid, km=c['km'], source='template', upload_id=up.id, entered_by=user.id))
        rate = VAT_RATES[c['vat']]
        for a, k in c['amounts'].items():
            g, n, v = split_vat(k, rate, c['gross'])
            db.add(Operation(ym=ym, vehicle_id=vid, article=a, amount_gross_kop=g, amount_net_kop=n, vat_kop=v, vat_rate=c['vat'],
                             vat_mode='gross' if c['gross'] else 'net', source='template', document=c['doc'], upload_id=up.id, entered_by=user.id))
    for sh in plan['shared']:
        old = list(db.scalars(select(Operation).where(Operation.ym == ym, Operation.vehicle_id.is_(None), Operation.article == sh['article'],
                                                      Operation.source == 'template')))
        backup['operations'] += [row_to_dict(o) for o in old]
        for o in old:
            db.delete(o)
        db.flush()
        g, n, v = split_vat(sh['kop'], VAT_RATES[sh['vat']], sh['gross'])
        db.add(Operation(ym=ym, vehicle_id=None, article=sh['article'], amount_gross_kop=g, amount_net_kop=n, vat_kop=v, vat_rate=sh['vat'],
                         vat_mode='gross' if sh['gross'] else 'net', source='template', document=sh['doc'], upload_id=up.id, entered_by=user.id))
    out = [plural(len(plan['cars']), 'машина', 'машины', 'машин')]
    if plan['shared']:
        out.append(plural(len(plan['shared']), 'общая статья', 'общие статьи', 'общих статей'))
    return [f'{ym_label(ym)}: ' + ', '.join(out)]


# ================================================================ таблица прибыли
def plan_profit(ctx: Ctx, wb) -> dict:
    cars, shared, bench = parsing.parse_profit(wb)
    year = parsing.parse_profit_year(wb)
    months = [f'{year}-{i:02d}' for i in range(1, 13)]
    db = ctx.db
    rows, notes, plan_months = [], [], []
    errors_cars = []
    known = {}
    for key, c in cars.items():
        v = ctx.vehicles.get(key)
        if v is None:
            errors_cars.append(c['label'])
        else:
            known[key] = v
    for i, ym in enumerate(months):
        vals = {k: c['p'][i] for k, c in cars.items() if k in known and c['p'][i] is not None}
        gen = (shared.get('gen') or [None] * 12)[i]
        if gen is None and not any(vals.values()):
            continue        # в файле нули-заготовки или пусто — месяц ещё не заполнен
        car_sum = sum(vals.values())
        total = (shared.get('total') or [None] * 12)[i]
        shared_sum = sum((shared.get(f) or [None] * 12)[i] or 0 for f in ('tractor', 'gen', 'service', 'amort'))
        disp = [ym_label(ym), str(len(vals)), rub(int(car_sum * 100)), rub(int(shared_sum * 100)), rub(int(total * 100)) if total is not None else '—']
        if ctx.is_open(ym):
            plan_months.append({'ym': ym, 'cars': {known[k].id: int(round(v * 100)) for k, v in vals.items()},
                                'shared': {f: int(round(((shared.get(f) or [None] * 12)[i] or 0) * 100)) for f in ('tractor', 'gen', 'service', 'amort')
                                           if (shared.get(f) or [None] * 12)[i] is not None},
                                'total': total, 'bench': bench.get(i + 1)})
            rows.append(row(disp, 'ok', 'будет записано'))
        else:
            cur = db.scalar(select(ReferenceValue.value).where(ReferenceValue.ym == ym, ReferenceValue.metric == 'fleet_total'))
            same = cur is not None and total is not None and abs(float(cur) - float(total)) < 0.01
            st = 'закрыт' if ctx.periods.get(ym) == 'closed' else 'не открыт'
            note = f'месяц {st} — {"совпадает с программой" if same else ("в программе " + rub(int(float(cur) * 100)) if cur is not None else "в программе нет данных")}'
            rows.append(row(disp, 'skip' if same or cur is None else 'warn_skip', note))
    for r in rows:
        if r['status'] == 'warn_skip':
            r['status'] = 'skip'
            notes.append(f'{r["cells"][0]}: в файле итог {r["cells"][4]}, а в программе другой. Месяц закрыт — чтобы обновить, владелец переоткрывает его.')
    if errors_cars:
        notes.append('В справочнике нет машин: ' + ', '.join(errors_cars) + '. Их прибыль не будет записана — добавьте машины в справочник и загрузите файл снова.')
    return {'kind': 'profit', 'months': [m['ym'] for m in plan_months], 'sections': [
        section('profit', f'Прибыль по машинам · {year}', ['Месяц', 'Машин', 'Прибыль машин', 'Общие и амортизация', 'Итог «Общий»'], rows, notes,
                f'{len(plan_months)} мес. будет записано')],
        'plan': {'months': plan_months}}


def apply_profit(ctx: Ctx, up: Upload, user: User, plan: dict, backup: dict) -> list[str]:
    db = ctx.db
    arts = {'tractor': 'shared_tractor', 'gen': 'shared_gen', 'service': 'shared_service', 'amort': 'shared_amort'}
    done = []
    for m in plan['months']:
        ym = m['ym']
        old = list(db.scalars(select(Operation).where(Operation.ym == ym, Operation.source.in_(('import', 'upload')),
                                                      Operation.article.in_(['profit_total', *arts.values()]))))
        backup['operations'] += [row_to_dict(o) for o in old]
        for o in old:
            db.delete(o)
        oldref = list(db.scalars(select(ReferenceValue).where(ReferenceValue.ym == ym, ReferenceValue.metric.in_(('fleet_total', 'cost_per_km', 'rate_per_km')))))
        backup['reference_values'] += [row_to_dict(r) for r in oldref]
        for r in oldref:
            db.delete(r)
        db.flush()
        for vid, kop in m['cars'].items():
            db.add(Operation(ym=ym, vehicle_id=int(vid), article='profit_total', amount_gross_kop=kop, amount_net_kop=kop, vat_kop=0,
                             vat_mode='unknown', source='upload', upload_id=up.id, entered_by=user.id))
        for f, kop in m['shared'].items():
            db.add(Operation(ym=ym, vehicle_id=None, article=arts[f], amount_gross_kop=kop, amount_net_kop=kop, vat_kop=0,
                             vat_mode='unknown', source='upload', upload_id=up.id, entered_by=user.id))
        created = []
        if m['total'] is not None:
            db.add(ReferenceValue(ym=ym, metric='fleet_total', value=Decimal(str(m['total'])), note=f'загрузка №{up.id}'))
            created.append([ym, 'fleet_total'])
        b = m.get('bench') or {}
        if 'cost' in b:
            db.add(ReferenceValue(ym=ym, metric='cost_per_km', value=Decimal(str(b['cost'])), note='налоговая' if b.get('tax') else None))
            created.append([ym, 'cost_per_km'])
        if 'rate' in b:
            db.add(ReferenceValue(ym=ym, metric='rate_per_km', value=Decimal(str(b['rate']))))
            created.append([ym, 'rate_per_km'])
        backup['reference_created'] += created
        done.append(f'{ym_label(ym)}: прибыль {plural(len(m["cars"]), "машины", "машин", "машин")}')
    return done


# ================================================================ файл парка
def plan_park(ctx: Ctx, wb) -> dict:
    parsing.review.clear()
    sections, plan, months = [], {}, set()
    names = wb.sheetnames

    # --- справочник «АВТО»
    if 'АВТО' in names:
        avto = parsing.parse_avto(wb)
        letter = {'М': 'MAN', 'Д': 'DAF', 'Ф': 'FAW'}
        rows, changes = [], []
        fields = {'plate': 'Госномер', 'model': 'Модель', 'make': 'Марка', 'year': 'Год', 'trailer': 'Прицеп',
                  'trailer_model': 'Модель прицепа', 'trailer_year': 'Год прицепа'}
        plates_taken = {v.plate: v.key for v in ctx.vehicles.values() if v.plate}
        for base, entries in sorted(avto.items(), key=lambda x: int(x[0])):
            for e in entries:
                cands = [v for k, v in ctx.vehicles.items() if re.fullmatch(rf'{base}[МДФ]?', k)]
                v = cands[0] if len(cands) == 1 and (cands[0].make in (None, e['make']) or not re.search('[МДФ]$', cands[0].key)) else \
                    next((c for c in cands if letter.get(c.key[-1]) == e['make']), None)
                if v is None:
                    key = base if base not in parsing.TWINS else base + {'MAN': 'М', 'DAF': 'Д', 'FAW': 'Ф'}.get(e['make'], '')
                    if key in ctx.vehicles:
                        rows.append(row([parsing.display(key), e['plate'], e['model'], 'ошибка'], 'error', 'не удалось однозначно сопоставить'))
                        continue
                    if e['plate'] in plates_taken:
                        rows.append(row([parsing.display(key), e['plate'], e['model'], ''], 'error', f'госномер уже у машины {parsing.display(plates_taken[e["plate"]])}'))
                        continue
                    changes.append({'create': key, 'name': parsing.display(key), 'odo': e.get('odo'),
                                    'set': {f: (int(e[f]) if f in ('year', 'trailer_year') and isinstance(e.get(f), float) else e.get(f)) for f in fields}})
                    plates_taken[e['plate']] = key
                    rows.append(row([parsing.display(key), e['plate'], e['model'], 'новая машина'], 'warn', 'будет добавлена в справочник'))
                    continue
                diff = {}
                for f in fields:
                    new, old = e.get(f), getattr(v, f)
                    if f in ('year', 'trailer_year') and isinstance(new, float):
                        new = int(new)
                    if new not in (None, '') and str(new) != str(old if old is not None else ''):
                        if f == 'plate' and new in plates_taken and plates_taken[new] != v.key:
                            rows.append(row([v.name, new, e['model'], ''], 'error', f'госномер уже у машины {parsing.display(plates_taken[new])}'))
                            diff = None
                            break
                        diff[f] = new
                if diff is None:
                    continue
                if diff:
                    changes.append({'id': v.id, 'name': v.name, 'set': diff, 'odo': e.get('odo')})
                    what = ', '.join(f'{fields[f].lower()}: {getattr(v, f) or "—"} → {val}' for f, val in diff.items())
                    rows.append(row([v.name, e['plate'], e['model'], what], 'warn' if v.edited else 'ok',
                                    ('машину меняли в программе — файл перезапишет: ' if v.edited else '') + 'будет обновлено'))
        unchanged = sum(len(x) for x in avto.values()) - len(rows)
        sections.append(section('avto', 'Справочник «АВТО»', ['Машина', 'Госномер', 'Модель', 'Изменения'], rows,
                                [plural(unchanged, 'машина совпадает', 'машины совпадают', 'машин совпадают') + ' со справочником — без изменений.'] if unchanged else [],
                                f'{len(changes)} изменений'))
        plan['vehicles'] = changes
        if changes:
            months.add('vehicles')

    # --- график работы
    sched = parsing.parse_schedule(wb)
    if sched:
        rows, plan_sched, closed = [], [], []
        for b in sched:
            ym = f"{b['y']}-{b['m']:02d}"
            if not ctx.is_open(ym):
                closed.append(ym)
                continue
            good, errs, new_cars = [], [], []
            for r in b['rows']:
                k = r['key']
                if k in parsing.TWINS:      # 308 без буквы — не угадываем, какая из двух машин
                    errs.append(r['raw'])
                    continue
                if k not in ctx.vehicles:
                    new_cars.append(k)
                good.append({'k': k, 'c': r['codes'], 'tx': r['texts'], 'tot': r['tot'], 'calc': r['calc'], 'row': r['row']})
            plan_sched.append({'ym': ym, 'rows': good, 'new': sorted(set(new_cars))})
            months.add(ym)
            st = 'error' if errs else ('warn' if new_cars else 'ok')
            note = '; '.join(filter(None, [
                plural(len(good), 'машина', 'машины', 'машин') + ' к записи',
                f'новые в справочнике: {", ".join(parsing.display(x) for x in sorted(set(new_cars)))}' if new_cars else '',
                f'не понял номер (двойник без буквы): {", ".join(errs)} — эти строки пропущу' if errs else '',
                'итоги по дням посчитаны по ячейкам (ручных итогов нет)' if not any(r['tot'] for r in b['rows']) else '']))
            rows.append(row([ym_label(ym), str(len(good)), str(sum((r['tot'] or r['calc']).get('worked', 0) for r in b['rows']))], st, note))
        notes = []
        if closed:
            notes.append(f'Закрытые месяцы не меняются: пропущено {len(closed)} ({ym_label(closed[0]).lower()} — {ym_label(closed[-1]).lower()}).')
        if not plan_sched:
            notes.append('Открытых месяцев в графике нет — ничего не будет записано.')
        sections.append(section('schedule', 'График работы', ['Месяц', 'Машин', 'Рабочих дней'], rows, notes, f'{len(plan_sched)} мес.'))
        plan['schedule'] = plan_sched
    # --- заправки
    if 'Анализ ЗАПРАВОК' in names:
        daily, keys = parsing.parse_fuel(wb)
        per_month = defaultdict(lambda: {'days': [], 'cars': set(), 'liters': 0.0})
        unknown = sorted(k for k in keys if k not in ctx.vehicles)
        for k, days in daily.items():
            if k not in ctx.vehicles:
                continue
            for d, (rn, e1, gl) in days.items():
                m = per_month[d[:7]]
                m['days'].append([k, d, rn, e1, gl])
                m['cars'].add(k)
                m['liters'] += rn + e1
        rows, plan_fuel, closed = [], [], []
        for ym in sorted(per_month):
            m = per_month[ym]
            disp = [ym_label(ym), str(len(m['cars'])), f"{m['liters']:,.0f}".replace(',', ' ') + ' л']
            if ctx.is_open(ym):
                plan_fuel.append({'ym': ym, 'days': m['days']})
                months.add(ym)
                rows.append(row(disp, 'ok', 'будет записано (заменит заправки месяца)'))
            else:
                closed.append(ym)
        notes = []
        if closed:
            notes.append(f'Закрытые месяцы не меняются: пропущено {len(closed)}.')
        if unknown:
            notes.append('Нет в справочнике, пропущено: ' + ', '.join(parsing.display(k) for k in unknown) + '.')
        sections.append(section('fuel', 'Заправки', ['Месяц', 'Машин', 'По картам'], rows, notes, f'{len(plan_fuel)} мес.'))
        plan['fuel'] = plan_fuel
    review = [r for r in parsing.review if any(ym_label(ym).lower() in r['where'].lower() for ym in months if ym != 'vehicles')]
    plan['review'] = review
    if not sections:
        raise UploadError('В файле нет ни графика работы, ни «Анализа ЗАПРАВОК», ни листа «АВТО».')
    return {'kind': 'park', 'months': sorted(months), 'sections': sections, 'plan': plan}


def apply_park(ctx: Ctx, up: Upload, user: User, plan: dict, backup: dict) -> list[str]:
    db = ctx.db
    done = []
    for ch in plan.get('vehicles', []):
        if 'create' in ch:
            v = Vehicle(key=ch['create'], name=ch['name'], is_active=True, manual=False)
            for f, val in ch['set'].items():
                setattr(v, f, val)
            v.odo_km = ch.get('odo')
            db.add(v)
            db.flush()
            ctx.vehicles[v.key] = v
            backup['created_vehicles'].append(v.id)
            log(db, user, 'vehicle', v.name, f'Машина добавлена (загрузка №{up.id})', None, v.plate or '')
        else:
            v = db.get(Vehicle, ch['id'])
            backup['vehicles'].append({'id': v.id, 'before': {f: getattr(v, f) for f in ch['set']}, 'after': ch['set']})
            for f, val in ch['set'].items():
                log(db, user, 'vehicle', v.name, f'{f} (загрузка №{up.id})', getattr(v, f), val)
                setattr(v, f, val)
            if ch.get('odo'):
                v.odo_km = ch['odo']
    if plan.get('vehicles'):
        done.append('справочник: ' + plural(len(plan['vehicles']), 'изменение', 'изменения', 'изменений'))
    for m in plan.get('schedule', []):
        for k in m['new']:
            if k not in ctx.vehicles:
                v = Vehicle(key=k, name=parsing.display(k), is_active=True)
                db.add(v)
                db.flush()
                ctx.vehicles[k] = v
                backup['created_vehicles'].append(v.id)
                db.add(ReviewItem(kind='Машина без госномера', value=v.name, where_text=f'график, {ym_label(m["ym"]).lower()}',
                                  action='добавлена из графика — укажите госномер в справочнике', upload_id=up.id))
        old = list(db.scalars(select(ScheduleMonth).where(ScheduleMonth.ym == m['ym'])))
        backup['schedule_months'] += [row_to_dict(o) for o in old]
        for o in old:
            db.delete(o)
        db.flush()
        for r in m['rows']:
            db.add(ScheduleMonth(ym=m['ym'], vehicle_id=ctx.vehicles[r['k']].id, codes=r['c'], texts=r['tx'], totals=r['tot'],
                                 calc=r['calc'], source_row=r['row'], upload_id=up.id))
        done.append(f'график {ym_label(m["ym"]).lower()}: ' + plural(len(m['rows']), 'машина', 'машины', 'машин'))
    for m in plan.get('fuel', []):
        y, mo = map(int, m['ym'].split('-'))
        start, end = dt.date(y, mo, 1), dt.date(y, mo, days_in(m['ym']))
        old = list(db.scalars(select(FuelDay).where(FuelDay.date >= start, FuelDay.date <= end)))
        backup['fuel_days'] += [row_to_dict(o) for o in old]
        for o in old:
            db.delete(o)
        db.flush()
        for k, d, rn, e1, gl in m['days']:
            db.add(FuelDay(vehicle_id=ctx.vehicles[k].id, date=dt.date.fromisoformat(d), rn_l=Decimal(str(rn)), e100_l=Decimal(str(e1)),
                           glonass_l=Decimal(str(gl)), upload_id=up.id))
        done.append(f'заправки {ym_label(m["ym"]).lower()}: ' + plural(len(m['days']), 'запись', 'записи', 'записей'))
    for r in plan.get('review', []):
        db.add(ReviewItem(kind=r['kind'], value=r['value'], where_text=r['where'], action=r['action'], upload_id=up.id))
    return done


# ================================================================ точки входа
PLANNERS = {'template': plan_template, 'profit': plan_profit, 'park': plan_park}
APPLIERS = {'template': apply_template, 'profit': apply_profit, 'park': apply_park}
KIND_NAMES = {'template': 'Шаблон FiberTrack', 'profit': 'Таблица прибыли', 'park': 'Файл парка'}


def make_preview(db: Session, path: str, chosen_ym: str | None, role: str) -> dict:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        raise UploadError('Не получилось открыть файл как Excel (.xlsx). Если это старый .xls — откройте его в Excel и сохраните как .xlsx.')
    kind = detect(wb)
    if kind == 'profit' and role != 'admin':
        raise UploadError('Таблицу прибыли загружает владелец.')
    ctx = Ctx(db)
    res = plan_template(ctx, wb, chosen_ym) if kind == 'template' else PLANNERS[kind](ctx, wb)
    res['kind_name'] = KIND_NAMES[kind]
    return res


def apply_upload(db: Session, up: Upload, user: User) -> list[str]:
    ctx = Ctx(db)
    closed = [m for m in up.months if m != 'vehicles' and not ctx.is_open(m)]
    if closed:
        raise UploadError(f'{ym_label(closed[0])} уже закрыт — загрузите файл заново после того, как владелец откроет месяц.')
    backup = {'operations': [], 'mileage': [], 'schedule_months': [], 'fuel_days': [], 'reference_values': [], 'reference_created': [],
              'vehicles': [], 'created_vehicles': []}
    done = APPLIERS[up.kind](ctx, up, user, up.plan, backup)
    db.flush()
    backup['created_counts'] = {m.__tablename__: db.query(m).filter(m.upload_id == up.id).count() for m in TRACKED}
    up.backup, up.status, up.applied_by, up.applied_at = backup, 'applied', user.id, now()
    log(db, user, 'upload', f'№{up.id}', f'Применена загрузка «{up.filename}»', None, '; '.join(done) or 'без изменений')
    return done


def rollback_upload(db: Session, up: Upload, user: User) -> None:
    if up.status != 'applied':
        raise UploadError('Откатить можно только применённую загрузку.')
    ctx = Ctx(db)
    closed = [m for m in up.months if m != 'vehicles' and not ctx.is_open(m)]
    if closed:
        raise UploadError(f'{ym_label(closed[0])} закрыт. Чтобы откатить, сначала переоткройте месяц.')
    later = [u for u in db.scalars(select(Upload).where(Upload.id > up.id, Upload.status == 'applied')) if set(u.months) & set(up.months)]
    if later:
        raise UploadError(f'После неё применена загрузка №{later[0].id} («{later[0].filename}») по тем же данным. Сначала откатите её.')
    b = up.backup or {}
    counts = {m.__tablename__: db.query(m).filter(m.upload_id == up.id).count() for m in TRACKED}
    if counts != b.get('created_counts', counts):
        raise UploadError('После загрузки эти данные меняли вручную — откат вернул бы их неверно. Загрузите исправленный файл заново.')
    for m in TRACKED:
        db.execute(delete(m).where(m.upload_id == up.id))
    db.execute(delete(ReviewItem).where(ReviewItem.upload_id == up.id))
    for ym, metric in b.get('reference_created', []):
        db.execute(delete(ReferenceValue).where(ReferenceValue.ym == ym, ReferenceValue.metric == metric))
    db.flush()
    for model, key in ((Operation, 'operations'), (Mileage, 'mileage'), (ScheduleMonth, 'schedule_months'), (FuelDay, 'fuel_days'),
                       (ReferenceValue, 'reference_values')):
        for d in b.get(key, []):
            db.add(dict_to_row(model, d))
    for ch in b.get('vehicles', []):
        v = db.get(Vehicle, ch['id'])
        for f, val in ch['before'].items():
            if str(getattr(v, f)) == str(ch['after'][f]):   # поле потом не меняли вручную — возвращаем
                setattr(v, f, val)
    for vid in b.get('created_vehicles', []):
        v = db.get(Vehicle, vid)
        if v is not None:
            used = db.query(Operation).filter(Operation.vehicle_id == vid).count() + db.query(Mileage).filter(Mileage.vehicle_id == vid).count() \
                + db.query(ScheduleMonth).filter(ScheduleMonth.vehicle_id == vid).count() + db.query(FuelDay).filter(FuelDay.vehicle_id == vid).count()
            if used == 0:
                db.delete(v)
    up.status, up.rolled_back_by, up.rolled_back_at = 'rolled_back', user.id, now()
    log(db, user, 'upload', f'№{up.id}', f'Откат загрузки «{up.filename}»', 'применена', 'откачена')
