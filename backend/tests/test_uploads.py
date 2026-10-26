"""Загрузки: шаблон (без файлов заказчика), файл парка и таблица прибыли (на файлах заказчика, если указаны)."""
import io
import os

import openpyxl
import pytest
from sqlalchemy import func, select

from app.db import SessionLocal
from app.models import Mileage, Operation, ReviewItem, ScheduleMonth
from tests.conftest import login

PROFIT, PARK = os.environ.get('FT_TEST_PROFIT'), os.environ.get('FT_TEST_PARK')
real = pytest.mark.skipif(not (PROFIT and PARK), reason='не заданы пути к файлам заказчика')


def filled_template(client, ym='2026-10', **cells) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f'/api/template/{ym}').content))
    ws = wb.active
    for ref, val in cells.items():
        ws[ref] = val
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, data: bytes, name='шаблон.xlsx', ym=None):
    return client.post('/api/uploads', data={'ym': ym} if ym else {}, files={'file': (name, data)})


def count(model, **where):
    with SessionLocal() as db:
        q = select(func.count()).select_from(model)
        for k, v in where.items():
            q = q.where(getattr(model, k) == v)
        return db.scalar(q)


def test_template_lists_month_cars_and_round_trips():
    u = login('input')
    wb = openpyxl.load_workbook(io.BytesIO(u.get('/api/template/2026-10').content))
    ws = wb.active
    assert ws.title == 'Октябрь 2026' and ws['A3'].value == '885' and ws['B3'].value == 'Т885ВХ58'
    data = filled_template(u, C3=11200, D3='1 100 000,50', F3=300000, R3=20, S3='22%', T3='да')
    r = upload(u, data)
    assert r.status_code == 200, r.text
    p = r.json()
    rows = p['sections'][0]['rows']
    assert rows[0]['status'] == 'ok', rows[0]
    assert p['can_apply'] and p['months'] == ['2026-10']
    assert u.post(f"/api/uploads/{p['id']}/apply").status_code == 200
    with SessionLocal() as db:
        rev = db.scalar(select(Operation).where(Operation.article == 'revenue', Operation.source == 'template'))
        assert (rev.amount_gross_kop, rev.amount_net_kop, rev.vat_kop) == (110000050, 90163975, 19836075)  # 1 100 000,50 ₽ с НДС 22 %
    assert count(Mileage, km=11200, worked_days=20) == 1
    e = u.get('/api/bootstrap').json()['input']['2026-10']['885']
    assert e['km'] == 11200 and e['amounts']['revenue'] == 1100000.5 and e['source'] == 'template'


def test_template_errors_are_shown_and_not_applied():
    u = login('input')
    data = filled_template(u, C3='много', D3='сто', F3=-1, R3=40)
    p = upload(u, data).json()
    row = p['sections'][0]['rows'][0]
    assert row['status'] == 'error'
    assert 'пробег' in row['note'] and 'не число' in row['note'] and 'отрицательная' in row['note'] and 'дней' in row['note']
    assert not p['can_apply']


def own_template(formulas_cached=True) -> bytes:
    """Шаблон, собранный бухгалтером: без служебного листа, с формулами, общие расходы в столбцах статей."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Октябрь 2026'
    ws['A1'] = 'FiberTrack · шаблон ввода · Октябрь 2026 · заполните жёлтые ячейки, суммы в рублях'
    head = ['Машина', 'Госномер', 'Пробег, км', 'Выручка от перевозок', 'Прочий доход', 'Топливо', 'Ремонт и запчасти', 'ТО', 'Шины',
            'Зарплата водителей', 'Суточные', 'Платные дороги и «Платон»', 'Страхование', 'Лизинг / кредит', 'Мойка', 'Штрафы', 'Прочее',
            'Отработано дней', 'Ставка НДС', 'Суммы с НДС (да/нет)', 'Документ']
    for c, h in enumerate(head, 1):
        ws.cell(2, c, h)
    ws.append(['885', 'Т885ВХ58', 31000, 1155000, None, 306000, '=31624.95-27000+174700', None, None, 204026.23, 27000,
               '=12084.2+17520.45', None, 0, 0, '=3442.62+O3', None, 19, '22%', 'да'])
    ws.append([])
    ws.append(['Общие расходы (без машины)'])
    ws.append(['Общехоз', None, '=1308677.79-143861', *[None] * 15, '22%', 'да'])
    ws.append(['Сервис', None, 143861, *[None] * 15, '22%', 'да'])
    ws.append(['Трактор', None, 0, 0.01, None, 7815, 0, None, None, 67180, *[None] * 8, '22%', 'да'])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_own_template_with_formulas_and_shared_block():
    u = login('input')
    p = upload(u, own_template()).json()        # openpyxl не сохраняет результат формул — программа досчитывает сама
    cars, shared = p['sections'][0]['rows'], p['sections'][1]['rows']
    assert cars[0]['status'] == 'warn' and 'пробег' in cars[0]['note']        # больше обычного — предупреждение, но пишется
    assert [r['status'] for r in shared] == ['ok', 'ok', 'ok'] and p['can_apply'], shared
    assert u.post(f"/api/uploads/{p['id']}/apply").status_code == 200
    with SessionLocal() as db:
        repair = db.scalar(select(Operation.amount_gross_kop).where(Operation.article == 'repair'))
        fines = db.scalar(select(Operation.amount_gross_kop).where(Operation.article == 'fines'))
        shared_kop = dict(db.execute(select(Operation.article, Operation.amount_gross_kop).where(Operation.vehicle_id.is_(None))).all())
    assert repair == 17932495 and fines == 344262                               # 31624,95 − 27000 + 174700; 3442,62 + 0
    assert shared_kop == {'shared_gen': 116481679, 'shared_service': 14386100, 'shared_tractor': 7499499}   # трактор: расходы − доход
    assert count(Mileage, km=31000, worked_days=19) == 1
    a = login('admin').get('/api/bootstrap').json()
    assert a['wdays']['2026-10']['885'] == 19 and a['fin']['2026-10']['shared']['g']['shared_tractor'] == 74994.99


def test_formula_evaluator():
    from app.uploads import formula_value
    ws = openpyxl.Workbook().active
    ws['A1'], ws['A2'] = 10, '=A1*2'
    assert formula_value('=(A1+5)/3-A2', ws) == -15
    assert formula_value('=SUM(A1:A2)', ws) is None and formula_value('=9**9**9', ws) is None


def test_template_for_closed_month_writes_nothing():
    a = login('admin')
    data = filled_template(a, ym='2026-08', C3=1000)
    p = upload(a, data).json()
    assert not p['can_apply'] and 'закрыт' in ' '.join(p['sections'][0]['notes'])


def test_reupload_replaces_and_rollback_order():
    u, a = login('input'), login('admin')
    first = upload(u, filled_template(u, C3=1000, D3=500)).json()
    u.post(f"/api/uploads/{first['id']}/apply")
    second = upload(u, filled_template(u, C3=2000, D3=700)).json()
    u.post(f"/api/uploads/{second['id']}/apply")
    assert count(Operation, source='template', article='revenue') == 1        # заменилось, а не сложилось
    assert a.post(f"/api/uploads/{first['id']}/rollback").status_code == 409  # сначала более позднюю
    assert a.post(f"/api/uploads/{second['id']}/rollback").status_code == 200
    with SessionLocal() as db:
        assert db.scalar(select(Mileage.km)) == 1000                           # вернулось значение первой загрузки
    assert a.post(f"/api/uploads/{first['id']}/rollback").status_code == 200
    assert count(Operation, source='template') == 0 and count(Mileage) == 0


def test_rollback_refused_after_manual_change():
    u, a = login('input'), login('admin')
    p = upload(u, filled_template(u, C3=1000, D3=500)).json()
    u.post(f"/api/uploads/{p['id']}/apply")
    u.put('/api/input/2026-10/885', json={'km': 1500})
    r = a.post(f"/api/uploads/{p['id']}/rollback")
    assert r.status_code == 409 and 'вручную' in r.json()['detail']


def test_manual_entry_replaces_template_data():
    u = login('input')
    p = upload(u, filled_template(u, C3=1000, D3=500, H3=40)).json()
    u.post(f"/api/uploads/{p['id']}/apply")
    u.put('/api/input/2026-10/885', json={'km': 1000, 'amounts': {'revenue': 600}})
    assert count(Operation, vehicle_id=1, article='revenue') == 1 and count(Operation, article='maintenance') == 0


def test_bad_files_are_explained():
    u = login('input')
    assert 'xlsx' in u.post('/api/uploads', files={'file': ('a.csv', b'1,2')}).json()['detail']
    assert 'Excel' in u.post('/api/uploads', files={'file': ('a.xlsx', b'not excel')}).json()['detail']
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'что-то своё'
    buf = io.BytesIO()
    wb.save(buf)
    assert 'Не узнал формат' in upload(u, buf.getvalue(), 'x.xlsx').json()['detail']


@real
def test_park_file_updates_open_months_only_and_rolls_back():
    from app.importer import import_history
    with SessionLocal() as db:
        import_history(db, PROFIT, PARK)
    before = count(ScheduleMonth)
    u, a = login('input'), login('admin')
    p = upload(u, open(PARK, 'rb').read(), 'парк.xlsx').json()
    titles = [s['title'] for s in p['sections']]
    assert titles == ['Справочник «АВТО»', 'График работы', 'Заправки']
    assert set(p['months']) <= {'2026-09', '2026-10', 'vehicles'}             # закрытые месяцы не трогаем
    assert p['sections'][0]['stats']['apply'] == 0                            # «АВТО» совпадает со справочником
    r = u.post(f"/api/uploads/{p['id']}/apply")
    assert r.status_code == 200 and count(ScheduleMonth) == before
    assert a.post(f"/api/uploads/{p['id']}/rollback").status_code == 200
    assert count(ScheduleMonth) == before and count(ReviewItem, upload_id=p['id']) == 0


@real
def test_profit_table_owner_only_and_closed_months_untouched():
    from app.importer import import_history
    with SessionLocal() as db:
        import_history(db, PROFIT, PARK)
    assert login('input').post('/api/uploads', files={'file': ('p.xlsx', open(PROFIT, 'rb'))}).status_code == 422
    p = login('admin').post('/api/uploads', files={'file': ('p.xlsx', open(PROFIT, 'rb'))}).json()
    assert p['months'] == [] and not p['can_apply']
    assert all('совпадает' in r['note'] for r in p['sections'][0]['rows'])


def test_entered_money_reaches_reports_but_not_the_accountant():
    u = login('input')
    p = upload(u, filled_template(u, C3=10000, D3=1220000, F3=366000)).json()
    u.post(f"/api/uploads/{p['id']}/apply")
    d = login('viewer').get('/api/bootstrap').json()
    car = d['fin']['2026-10']['cars']['885']
    assert car['g'] == {'revenue': 1220000, 'fuel': 366000} and car['n'] == {'revenue': 1000000, 'fuel': 300000}
    assert d['km']['2026-10']['885'] == 10000
    b = u.get('/api/bootstrap').json()
    assert b['fin'] == {} and b['km'] == {}



def test_template_alone_starts_month_and_adds_new_cars():
    """Работа только через шаблон: новой машины нет в справочнике, месяц ещё не заводили."""
    wb = openpyxl.load_workbook(io.BytesIO(own_template()))
    ws = wb.active
    ws.title, ws['A1'] = 'Сентябрь 2026', 'FiberTrack · шаблон ввода · Сентябрь 2026'
    ws['A3'], ws['B3'] = '777', 'а777аа 58'
    buf = io.BytesIO()
    wb.save(buf)
    u, a = login('input'), login('admin')
    p = upload(u, buf.getvalue()).json()
    assert p['months'] == ['2026-09'] and p['can_apply'], p
    assert 'новая машина' in p['sections'][0]['rows'][0]['note'] and any('ещё не заведён' in n for n in p['sections'][0]['notes'])
    assert u.post(f"/api/uploads/{p['id']}/apply").status_code == 200
    d = a.get('/api/bootstrap').json()
    assert d['periods']['2026-09'] == 'open'
    assert d['vehicles']['777']['name'] == '777' and d['vehicles']['777']['plate'] == 'А777АА58'
    assert d['fin']['2026-09']['cars']['777']['g']['revenue'] == 1155000
    assert a.post(f"/api/uploads/{p['id']}/rollback").status_code == 200
    assert '777' not in a.get('/api/bootstrap').json()['vehicles']            # откат убирает и добавленную машину


def test_template_does_not_reopen_closed_month():
    wb = openpyxl.load_workbook(io.BytesIO(own_template()))
    wb.active['A1'] = 'FiberTrack · шаблон ввода · Август 2026'
    buf = io.BytesIO()
    wb.save(buf)
    p = upload(login('admin'), buf.getvalue()).json()
    assert not p['can_apply'] and 'закрыт' in ' '.join(p['sections'][0]['notes'])
